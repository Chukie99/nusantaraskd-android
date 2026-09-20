import json
from docx import Document

json_path = "F:/hermes/nusantaraskd-android/app/src/main/assets/questions_bank_1100.json"
docx_path = "F:/hermes/nusantaraskd-android/Bank_Soal_SKD_CPNS_1100.docx"

print("Membaca file JSON...")
with open(json_path, "r", encoding="utf-8") as f:
    data = json.load(f)

print(f"Total soal di JSON: {len(data)}")

doc = Document()
doc.add_heading("BANK SOAL RESMI SKD CPNS 2025/2026", 0)
doc.add_paragraph("Dokumen ini berisi 1100 soal lengkap (TWK, TIU, TKP) beserta kunci jawaban dan pembahasan untuk aplikasi Nusantara SKD Android.")

for idx, q in enumerate(data, 1):
    doc.add_heading(f"Soal {idx}: [{q['category']}] ({q['id']})", level=2)
    doc.add_paragraph(q['questionText'])
    
    doc.add_paragraph(f"A. {q['optionA']}")
    doc.add_paragraph(f"B. {q['optionB']}")
    doc.add_paragraph(f"C. {q['optionC']}")
    doc.add_paragraph(f"D. {q['optionD']}")
    doc.add_paragraph(f"E. {q['optionE']}")
    
    if q['category'] == 'TKP':
        doc.add_paragraph(f"Kunci/Bobot: A({q['weightA']}), B({q['weightB']}), C({q['weightC']}), D({q['weightD']}), E({q['weightE']})")
    else:
        doc.add_paragraph(f"Kunci Jawaban: {q['correctAnswer']}")
        
    doc.add_paragraph(f"Pembahasan: {q['explanation']}")
    doc.add_paragraph("-" * 40)

print(f"Menyimpan ke file DOCX: {docx_path}")
doc.save(docx_path)
print("Berhasil membuat file DOCX!")
